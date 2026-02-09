from openpyxl import Workbook
from programs.models.ifashe_models import (
    Family,
    SchoolSupport,
    SponsoredChild,
    Parent,
    DressingDistribution,
    Sponsorship,
)
from utils.reports.ifashe.base import BasePDFReport


class IFASHESummaryPDFReport(BasePDFReport):
    title = "IFASHE – Program Summary"

    def generate(self):
        self.add_title()

        rows = [
            [
                Family.objects.count(),
                Parent.objects.count(),
                SponsoredChild.objects.count(),
                DressingDistribution.objects.count(),
            ]
        ]

        self.add_table(
            headers=[
                "Families",
                "Parents",
                "Children",
                "Clothes Distributed",
            ],
            rows=rows,
        )
        self.build()


class IfasheSummaryExcelReport:
    def generate(self, file_path):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("IFASHE Summary")
        ws.append(["Metric", "Value"])
        ws.append(["Total Families", Family.objects.count()])
        ws.append(["Total Parents", Parent.objects.count()])
        ws.append(["Total Children", SponsoredChild.objects.count()])
        ws.append(
            ["Active Sponsorships", Sponsorship.objects.filter(status="active").count()]
        )
        ws.append(
            [
                "Pending School Supports",
                SchoolSupport.objects.filter(payment_status="pending").count(),
            ]
        )

        wb.save(file_path)
