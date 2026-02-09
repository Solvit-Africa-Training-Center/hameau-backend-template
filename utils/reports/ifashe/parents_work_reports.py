from openpyxl import Workbook
from programs.models.ifashe_models import Parent, ParentWorkContract, ParentAttendance
from utils.reports.ifashe.base import BasePDFReport


class ParentWorkCompliancePDFReport(BasePDFReport):
    title = "IFASHE – Parent Work Compliance"

    def generate(self):
        self.add_title()

        rows = []

        contracts = ParentWorkContract.objects.select_related("parent").iterator(200)

        for contract in contracts:
            attendance_count = ParentAttendance.objects.filter(
                work_record=contract, status=ParentAttendance.PRESENT
            ).count()

            rows.append(
                [
                    contract.parent.full_name,
                    contract.job_role,
                    contract.status,
                    attendance_count,
                ]
            )

        self.add_table(
            headers=[
                "Parent",
                "Job Role",
                "Contract Status",
                "Days Present",
            ],
            rows=rows,
        )

        self.build()


class ParentWorkExcelReport:
    def generate(self, file_path):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Parent Work Report")

        ws.append(
            [
                "Parent Name",
                "Family",
                "Job Role",
                "Attendance Rate",
                "Performance Score",
                "Compliance Status",
            ]
        )

        qs = Parent.objects.select_related("family")

        for parent in qs.iterator(200):
            ws.append(
                [
                    parent.full_name,
                    parent.family.family_name if parent.family else "",
                    getattr(parent, "job_role", ""),
                    getattr(parent, "attendance_rate", ""),
                    getattr(parent, "performance_score", ""),
                    getattr(parent, "compliance_status", ""),
                ]
            )

        wb.save(file_path)
