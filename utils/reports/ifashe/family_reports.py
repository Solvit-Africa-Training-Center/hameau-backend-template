from programs.models.ifashe_models import Family
from utils.reports.ifashe.base import BasePDFReport
from openpyxl import Workbook


class FamilyOverviewPDFReport(BasePDFReport):
    title = "IFASHE – Families Overview"

    def generate(self):
        self.add_title()

        rows = []
        qs = Family.objects.prefetch_related("parents", "children").order_by(
            "family_name"
        )

        for family in qs.iterator(200):
            rows.append(
                [
                    family.family_name,
                    family.province,
                    family.vulnerability_level,
                    family.parents.count(),
                    family.children.count(),
                ]
            )

        self.add_table(
            headers=[
                "Family",
                "Province",
                "Vulnerability",
                "Parents",
                "Children",
            ],
            rows=rows,
        )
        self.build()


class FamilyOverviewExcelReport:
    def generate(self, file_path):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Families Overview")

        ws.append(
            [
                "Family Name",
                "Province",
                "Vulnerability Level",
                "Parents Count",
                "Children Count",
            ]
        )

        qs = Family.objects.prefetch_related("parents", "children").order_by(
            "family_name"
        )

        for family in qs.iterator(200):
            ws.append(
                [
                    family.family_name,
                    family.province,
                    family.vulnerability_level,
                    family.parents.count(),
                    family.children.count(),
                ]
            )

        wb.save(file_path)
