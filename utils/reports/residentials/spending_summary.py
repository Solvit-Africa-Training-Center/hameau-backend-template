from programs.models.residentials_models import (
    HealthRecord,
    ChildEducation,
    ChildInsurance,
    ResidentialFinancialPlan,
    Child,
)
from utils.reports.residentials.base import BasePDFReport
from openpyxl import Workbook
from decimal import Decimal
from django.db.models import Sum, Q


class SpendingSummaryMixin:
    def __init__(self, request=None):
        self.request = request

    def get_date_range(self, obj):
        request = self.request
        if not request:
            raise ValueError("Request is required to fetch the date range")

        date_from = request.query_params.get("date_from")
        date_to = request.query_params.get("date_to")

        if not date_from or not date_to:
            raise ValueError(
                "Both 'date_from' and 'date_to' query parameters are required"
            )

        return {
            "from": date_from,
            "to": date_to,
        }

    def get_date_filters(self, date_field):
        filters = {}
        if self.request:
            start_date = self.request.query_params.get(
                "start_date"
            ) or self.request.query_params.get("date_from")
            end_date = self.request.query_params.get(
                "end_date"
            ) or self.request.query_params.get("date_to")

            if start_date:
                filters[f"{date_field}__gte"] = start_date
            if end_date:
                filters[f"{date_field}__lte"] = end_date
        return filters

    def _calculate_total_costs(self, children_queryset):
        date_range = self.get_date_range(None)
        start_date = date_range["from"]
        end_date = date_range["to"]

        health_cost = HealthRecord.objects.filter(
            child__in=children_queryset,
            visit_date__gte=start_date,
            visit_date__lte=end_date,
        ).aggregate(total=Sum("cost"))["total"] or Decimal("0.00")

        edu_cost = ChildEducation.objects.filter(
            child__in=children_queryset,
            start_date__gte=start_date,
            start_date__lte=end_date,
        ).aggregate(total=Sum("cost"))["total"] or Decimal("0.00")

        ins_cost = ChildInsurance.objects.filter(
            child__in=children_queryset,
            start_date__gte=start_date,
            start_date__lte=end_date,
        ).aggregate(total=Sum("cost"))["total"] or Decimal("0.00")

        food_cost = ResidentialFinancialPlan.objects.filter(
            child__in=children_queryset, month__gte=start_date, month__lte=end_date
        ).aggregate(total=Sum("food_cost"))["total"] or Decimal("0.00")

        return health_cost + edu_cost + ins_cost + food_cost

    def get_normal_spending(self):
        children = Child.objects.filter(
            Q(special_needs__isnull=True) | Q(special_needs__exact="")
        )
        return self._calculate_total_costs(children)

    def get_special_diet_spending(self):
        children = Child.objects.exclude(
            Q(special_needs__isnull=True) | Q(special_needs__exact="")
        )
        return self._calculate_total_costs(children)

    def get_education_spending(self):
        return ChildEducation.objects.filter(
            **self.get_date_filters("start_date")
        ).aggregate(total=Sum("cost"))["total"] or Decimal("0.00")

    def get_total_spending(self):
        return self.get_normal_spending() + self.get_special_diet_spending()


class SpendingSummaryPDFReport(SpendingSummaryMixin, BasePDFReport):
    title = "Residential Spending Summary"

    def __init__(self, file_path, request=None):
        SpendingSummaryMixin.__init__(self, request=request)
        BasePDFReport.__init__(self, file_path)

    def generate(self):
        try:
            date_range = self.get_date_range(None)
            start_date = date_range["from"]
            end_date = date_range["to"]

            self.add_title()

            rows = [
                ("Normal Spending", self.get_normal_spending()),
                ("", ""),
                ("Special Diet Spending", self.get_special_diet_spending()),
                ("", ""),
                ("Education Spending", self.get_education_spending()),
                ("", ""),
                ("Total Spending", self.get_total_spending()),
            ]

            self.add_table(headers=["Category", "Amount"], rows=rows)
            self.build()

        except ValueError as e:
            self.add_title()
            self.add_table(headers=["Error"], rows=[[str(e)]])
            self.build()


class SpendingSummaryExcelReport(SpendingSummaryMixin):
    def __init__(self, request=None):
        super().__init__(request=request)

    def generate(self, file_path):
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("Spending Summary")

        ws.append(["Category", "Amount"])

        ws.append(["Normal Spending", f"{self.get_normal_spending():.2f}"])
        ws.append(["Special Diet Spending", f"{self.get_special_diet_spending():.2f}"])
        ws.append(["Education Spending", f"{self.get_education_spending():.2f}"])
        ws.append(["Total Spending", f"{self.get_total_spending():.2f}"])

        wb.save(file_path)
